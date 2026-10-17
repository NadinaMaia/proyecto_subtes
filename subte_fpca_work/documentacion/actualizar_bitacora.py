from __future__ import annotations

from datetime import date
from pathlib import Path
import json

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "subte_fpca_work"
OUT = WORK / "resultados"
DOC_PATH = WORK / "documentacion" / "Bitacora_Tecnica_Proyecto_Subte_FPCA.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172A3A"
MUTED = "5C6770"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
GOLD = "FFF4CE"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_run(run, size=11, bold=False, color=INK, italic=False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", *, bold=False, color=INK, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    style_run(r, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc, text: str, level=0) -> None:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 + 0.25 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    style_run(p.add_run(text))


def add_number(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    style_run(p.add_run(text))


def add_table(doc, headers, rows, widths_dxa, *, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        style_run(p.add_run(str(header)), size=font_size, bold=True, color=DARK_BLUE)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)

    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cell = cells[j]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(str(value)), size=font_size)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(paragraph.add_run("Página "), size=9, color=MUTED)
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    settings = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def read_optional_csv(name: str) -> pd.DataFrame:
    path = OUT / name
    return pd.read_csv(path) if path.exists() else pd.DataFrame()


def build() -> None:
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "Bitácora técnica - Proyecto Subte FPCA"
    doc.core_properties.subject = "Procesamiento histórico, FPCA, clustering y MaxVol"
    doc.core_properties.author = "Proyecto Subte FPCA"

    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "PROYECTO SUBTE FPCA  |  BITÁCORA TÉCNICA"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        style_run(run, size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    add_para(doc, "PROYECTO SUBTE FPCA", bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    style_run(p.add_run("Bitácora técnica y metodológica"), size=26, bold=True, color=INK)
    add_para(
        doc,
        "Procesamiento histórico de molinetes 2014-2026, auditoría de calidad, FPCA segmentada, clustering y selección de prototipos MaxVol",
        color=MUTED,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=18,
    )

    add_table(
        doc,
        ["Campo", "Valor"],
        [
            ["Notebook principal", "Proyecto_Subte_FPCA_2014-26.ipynb"],
            ["Notebook del dashboard", "creacion_dashboard.ipynb"],
            ["Carpeta de trabajo", "subte_fpca_work"],
            ["Versión de estaciones", "v2-aliases-inventario-y-mapa-2026-08-31"],
            ["Última actualización", date.today().isoformat()],
            ["Estado", "Análisis completo y dashboard interactivo validado; operación habitual documentada"],
        ],
        [2100, 7260],
        font_size=9.5,
    )

    add_para(
        doc,
        "Propósito de esta bitácora",
        bold=True,
        color=DARK_BLUE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=5,
    )
    add_para(
        doc,
        "Registrar decisiones, problemas encontrados, reglas de limpieza, ajustes de rendimiento, metodología analítica, validaciones y próximos pasos. Este documento debe actualizarse junto con cada cambio sustantivo del notebook.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    doc.add_page_break()

    doc.add_heading("1. Objetivo y alcance", level=1)
    add_para(
        doc,
        "El proyecto estudia la evolución del uso del Subte de Buenos Aires a partir de registros de molinetes en intervalos de 15 minutos. La unidad funcional es línea-estación por año y segmento de día. El análisis separa Laborable, Sábado y Domingo/Feriado para evitar que la diferencia entre calendarios domine los patrones entre estaciones.",
    )
    add_bullet(doc, "Período analizado: 2014-2026; 2026 se trata como año parcial mientras no cierre el calendario.")
    add_bullet(doc, "Universo: líneas A, B, C, D, E y H. El Premetro (PM) queda fuera del análisis.")
    add_bullet(doc, "Variable principal: PAX_TOTAL agregado por estación, fecha y slot de 15 minutos.")
    add_bullet(doc, "Propósito estadístico: describir formas horarias, clusters comparables en el tiempo, rupturas pandémicas, anomalías y prototipos geométricos.")
    add_bullet(doc, "Alcance del dashboard: operación habitual; las extensiones por eventos y los slots nocturnos aislados quedan fuera del día típico mostrado.")

    doc.add_heading("2. Arquitectura del procesamiento", level=1)
    add_table(
        doc,
        ["Capa", "Contenido", "Motivo"],
        [
            ["Fuentes", "CSV y miembros CSV dentro de ZIP", "Conservar los datos originales sin reescritura"],
            ["Estandarizados", "Un Parquet por fuente", "Lectura incremental y formatos históricos homogéneos"],
            ["Deduplicados", "Parquet anual solo cuando hay claves repetidas", "Resolver conflictos sin cargar todo en memoria"],
            ["Station-slot", "Estación-fecha-slot por año", "Reducir drásticamente el volumen para auditorías y FPCA"],
            ["Vista canónica", "Alias de línea/estación y descartes", "Corregir identidades sin releer 122 fuentes"],
            ["Resultados", "CSV, Markdown y ZIP", "Entregables reproducibles y auditables"],
        ],
        [1500, 3720, 4140],
        font_size=8.8,
    )

    doc.add_heading("3. Historial de problemas y tratamientos", level=1)
    doc.add_heading("3.1 Formatos de fecha", level=2)
    add_para(
        doc,
        "El archivo formato_fechas_por_archivo.csv es la fuente de verdad para el parseo. La regla explícita por archivo evita inferencias ambiguas, especialmente cuando día y mes son ambos menores o iguales a 12.",
    )
    add_table(
        doc,
        ["Caso", "Problema", "Tratamiento implementado"],
        [
            ["2014", "La fuente usa YYYY-MM-DD y un encabezado desplazado; inicialmente aparecía sin días válidos", "Ruta resuelta desde la raíz del proyecto, encabezado detectado y fecha parseada explícitamente"],
            ["2020", "Ene-Feb y Nov-Dic usan DD/MM/YYYY; Mar-Oct usa MM/DD/YYYY", "Regla mixed_2020_blocks con límites auditados por posición de fila"],
            ["2021", "Conviven DD/MM/YYYY y fechas ISO YYYY-MM-DD", "Regla mixed_iso_dmy: ISO primero y DD/MM/YYYY para el resto"],
            ["General", "La autodetección podía intercambiar día y mes", "Catálogo por archivo, firma de caché y reproceso solo de fuentes afectadas"],
        ],
        [1200, 3420, 4740],
        font_size=8.6,
    )

    doc.add_heading("3.2 Intervalos", level=2)
    add_para(
        doc,
        "La fuente oficial informa slots de 15 minutos. Para 2014, intervalos inclusivos como 09:15:00-09:29:00 se reconocen como 15 minutos. Se separan dos controles: duración objetivo y alineación a la grilla. Las duraciones no estándar quedan disponibles para inspección en intervalos_faltantes_muestra.csv y en la distribución de duraciones.",
    )

    doc.add_heading("3.3 Premetro", level=2)
    add_para(
        doc,
        "Los archivos o filas identificados como PM corresponden al Premetro y se excluyen del universo. La regla opera tanto por fuente como por código de línea para evitar que registros de prueba contaminen estaciones, intervalos o perfiles.",
    )

    doc.add_heading("3.4 Duplicados", level=2)
    add_bullet(doc, "Duplicado exacto: repetición de fecha, intervalo, línea, estación, molinete y todos los campos de pasajeros; se cuenta n-1.")
    add_bullet(doc, "Clave repetida: misma fecha, intervalo, línea, estación y molinete, aunque cambien los pasajeros.")
    add_bullet(doc, "Clave conflictiva: una clave repetida con diferencias en PAX_TOTAL o sus componentes.")
    add_bullet(doc, "Resolución: conservar una fila por clave, priorizando el mayor PAX_TOTAL y luego el origen más reciente según el orden definido.")

    doc.add_heading("3.5 Identidad de estaciones", level=2)
    add_para(
        doc,
        "El conteo inflado surgía porque estacion_id concatenaba texto crudo de línea y estación. La corrección mantiene los originales para auditoría, construye línea y estación canónicas y agrega nuevamente los station-slots que convergen en la misma identidad.",
    )
    add_table(
        doc,
        ["Tipo", "Ejemplos", "Acción"],
        [
            ["Código de línea", "LINEAA, LINEA_A, A", "Unificar como A; análogo para B, C, D, E y H"],
            ["Codificación", "Variantes corruptas de SAENZ PEÑA y AGÜERO", "Unificar por patrón controlado"],
            ["Sufijo", "PUEYRREDON.B, RETIRO.C, CALLAO.D, RETIRO.E", "Quitar sufijo cuando línea+nombre identifican la estación"],
            ["No informado", "#N/D, SIN ESTACION", "Excluir y registrar motivo"],
            ["Prueba", "PRUEBA | PRUEBA", "Excluir y registrar motivo"],
            ["Asignación imposible", "LORIA en B, C, D, E o H", "Excluir; Loria pertenece a la línea A"],
        ],
        [1500, 3930, 3930],
        font_size=8.7,
    )

    doc.add_heading("4. Validación del inventario de estaciones", level=1)
    add_para(
        doc,
        "La vista canónica fue contrastada directamente contra los Parquet agregados. El resultado separa estaciones observadas de estaciones esperadas en la red. En 2020 se observan 89 porque Plaza de Mayo no tiene registros en el archivo; esto es un faltante de cobertura y no una reducción del inventario nominal.",
    )
    station_rows = [
        [2014, 83, 83, "OK"], [2015, 85, 85, "OK"], [2016, 86, 86, "OK"],
        [2017, 86, 86, "OK"], [2018, 87, 87, "OK"], [2019, 90, 90, "OK"],
        [2020, 89, 90, "Plaza de Mayo sin observaciones"], [2021, 90, 90, "OK"],
        [2022, 90, 90, "OK"], [2023, 90, 90, "OK"], [2024, 90, 90, "OK"],
        [2025, 90, 90, "OK tras descartes"], [2026, 90, 90, "OK; año parcial"],
    ]
    add_table(doc, ["Año", "Observadas", "Esperadas", "Lectura"], station_rows, [1000, 1500, 1500, 5360], font_size=8.8)
    add_para(
        doc,
        "Descartes validados: diez registros station-slot de Loria asignados a líneas incorrectas en 2025, seis #N/D en 2025, un #N/D en 2026 y un registro PRUEBA en 2025. El detalle exacto se exporta a estaciones_descartadas.csv.",
        italic=True,
        color=MUTED,
    )

    doc.add_heading("5. Rendimiento, memoria y caché", level=1)
    add_bullet(doc, "Lectura por chunks de 300.000 filas y escritura Parquet con compresión Zstandard.")
    add_bullet(doc, "Manifiesto por fuente con tamaño, fecha de modificación, formato, años y conteos de duración; las fuentes sin cambios se reutilizan.")
    add_bullet(doc, "DuckDB en memoria para evitar bloqueos entre kernels; los resultados persistentes permanecen en Parquet.")
    add_bullet(doc, "Límite automático de memoria: aproximadamente 55% de RAM total y 75% de RAM disponible, tomando el mínimo.")
    add_bullet(doc, "Máximo automático de cuatro hilos y ajuste adicional según memoria disponible.")
    add_bullet(doc, "Directorio temporal explícito, reserva mínima de 3 GiB y límite de spill equivalente a 85% del espacio utilizable.")
    add_bullet(doc, "Capas anuales de deduplicación, station-slot, cobertura y auditoría: solo se reconstruyen años afectados.")
    add_bullet(doc, "La normalización de estaciones es una vista liviana; no invalida la caché de 122 fuentes.")

    doc.add_heading("6. Metodología estadística", level=1)
    doc.add_heading("6.1 Curvas y cobertura", level=2)
    add_para(doc, "Se agregan molinetes por estación-fecha-slot. Los slots esperados se estiman por año y segmento de día. Se descartan curvas con cobertura inferior al umbral y se conserva el volumen diario mediano como variable separada de la forma.")

    doc.add_heading("6.2 Suavizado y FPCA", level=2)
    add_para(doc, "Las curvas normalizadas se suavizan con base B-spline y penalización de segunda diferencia. Lambda se elige por GCV dentro de cada segmento. La FPCA se ajusta sobre años completos no pandémicos; 2020-2021 y años parciales se proyectan después sobre ese espacio estable.")

    doc.add_heading("6.3 Clustering", level=2)
    add_para(doc, "K-means se ajusta sobre los scores retenidos. El número de grupos se evalúa mediante silhouette, estabilidad ARI entre repeticiones y tamaño mínimo. Las etiquetas se ordenan por hora del pico del perfil mediano para facilitar su lectura temporal.")

    modelos = read_optional_csv("resumen_modelos_segmentados.csv")
    if not modelos.empty:
        rows = []
        for r in modelos.itertuples(index=False):
            rows.append([
                r.segmento_dia, int(r.curvas), int(r.componentes_90), int(r.k),
                f"{r.silhouette:.3f}", f"{r.ari_estabilidad:.3f}",
            ])
        add_para(doc, "Ejecución validada después de la normalización de estaciones:", bold=True, color=DARK_BLUE)
        add_table(doc, ["Segmento", "Curvas", "FPC 90%", "k", "Silhouette", "ARI"], rows, [2400, 1200, 1200, 700, 1700, 2160], font_size=8.8)

    doc.add_heading("6.4 MaxVol", level=2)
    add_para(
        doc,
        "MaxVol se aplica por segmento sobre los scores estandarizados de las FPC necesarias para alcanzar el objetivo de varianza. Se agrega una columna de unos para representar geometría afín. Una selección QR con pivoteo inicializa el algoritmo y los intercambios continúan hasta que el mayor coeficiente de interpolación queda por debajo de 1,05 o se alcanza el máximo de iteraciones.",
    )
    add_bullet(doc, "Los prototipos son observaciones reales y extremas; no son centroides.")
    add_bullet(doc, "Se muestran como estrellas sobre todos los scores FPCA y también como curvas horarias.")
    add_bullet(doc, "La tabla prototipos_maxvol.csv conserva estación, año, cluster, pico, volumen, referencia y diagnóstico numérico.")
    add_bullet(doc, "Los casos pandémicos o parciales pueden ser elegidos por su excepcionalidad; deben interpretarse con cobertura y contexto.")

    doc.add_heading("6.5 Preparación del dashboard geográfico", level=2)
    add_para(
        doc,
        "Se incorporó una capa local de 90 coordenadas canónicas proveniente del GeoJSON oficial "
        "de Buenos Aires Data. Los nombres extendidos del catálogo se traducen al diccionario "
        "histórico usado por el análisis. La unión se valida: una estación analítica sin coordenadas "
        "detiene la generación del insumo del mapa.",
    )
    add_bullet(doc, "El mapa representa un día típico mediante la mediana de pasajeros por estación, segmento y slot de 15 minutos.")
    add_bullet(doc, "El color identifica la línea; el tamaño del círculo expresa volumen absoluto y usa una escala estable entre horas para evitar amplificar volúmenes nocturnos mínimos.")
    add_bullet(doc, "El trazado se apoya en OpenStreetMap y conecta las estaciones oficiales en el orden ferroviario de cada línea; no requiere API key.")
    add_bullet(doc, "La capa de salida mapa_estaciones_perfil_tipico.csv está lista para una interfaz interactiva con año, segmento de día, líneas y deslizador horario.")
    add_bullet(doc, "El año 2026 se expone como lectura actual, pero debe estar rotulado como parcial; 2025 es la referencia anual completa sugerida.")

    doc.add_heading("6.6 Alcance horario y diseño del dashboard", level=2)
    add_para(
        doc,
        "El dashboard se diseñó para estudiar horarios corrientes y no anomalías operativas. Se adoptan ventanas explícitas: laborables 05:30-23:30, sábados 06:00-24:00 y domingos/feriados 08:00-22:00. Dentro de esas ventanas se conserva todo el dato disponible; sólo las franjas fuera de horario habitual se representan como cero.",
    )
    add_bullet(doc, "El deslizador se restringe automáticamente a la ventana del tipo de día seleccionado y no admite ingreso numérico manual.")
    add_bullet(doc, "El botón Reproducir/Pausar avanza en intervalos de 15 minutos. Durante la reproducción sólo se recalcula el mapa; al pausar, la hora se marca en los gráficos con eje temporal.")
    add_bullet(doc, "El linking and brushing mantiene una selección independiente por gráfico y aplica la intersección de las selecciones activas a todo el panel.")
    add_bullet(doc, "Reset view/Reset axes elimina únicamente la selección originada en ese gráfico; el botón global elimina todas las selecciones.")
    add_bullet(doc, "Los redibujados automáticos con selectedData=None no borran filtros persistentes, evitando el retorno momentáneo al estado original.")
    add_bullet(doc, "El mapa conserva el color de línea en el anillo del círculo y expresa la intensidad temporal con un relleno secuencial; el nombre oficial y las métricas aparecen en el hover.")
    add_bullet(doc, "La identidad visual incorpora el logo Demanda Subte, amarillo y negro, y distintivos circulares de línea sin borde blanco.")
    add_bullet(doc, "La escala '% del día' se reserva para perfiles temporales: expresa la proporción del tráfico diario típico concentrada en cada slot, no pasajeros absolutos.")
    add_bullet(doc, "En duración del pico, el largo de la barra es la racha continua con demanda al menos igual al 80% del máximo puntual; el color es la proporción del día concentrada en la mejor ventana de 60 minutos.")
    add_bullet(doc, "El ranking comparativo de picos separa las estaciones con meseta más extensa de aquellas con más pasajeros en su mejor ventana de 60 minutos.")

    doc.add_heading("7. Interpretación analítica disponible", level=1)
    clusters = read_optional_csv("caracterizacion_clusters.csv")
    if not clusters.empty:
        add_para(doc, "Los siguientes patrones corresponden a la ejecución completa validada con estaciones canónicas:", italic=True, color=MUTED)
        for r in clusters.itertuples(index=False):
            if r.balance_manana_tarde > 0.05:
                balance = "mayor peso matutino"
            elif r.balance_manana_tarde < -0.20:
                balance = "mayor peso vespertino"
            else:
                balance = "perfil relativamente equilibrado"
            add_bullet(
                doc,
                f"{r.cluster_id} ({r.segmento_dia}): pico {r.pico}, {balance}, "
                f"mediana diaria {r.pax_diarios_mediana:,.0f} y silhouette mediana {r.silhouette_mediana:.3f}.",
            )
    add_para(
        doc,
        "Lectura sustantiva esperada: en laborables se separan perfiles con fuerte entrada matutina de perfiles con concentración de salida por la tarde. Los sábados aparecen formas de mediodía más o menos concentradas y un grupo más vespertino. Domingos y feriados se desplazan hacia la tarde-noche. Esta interpretación describe forma relativa; no implica que todas las estaciones del grupo tengan el mismo volumen.",
    )

    doc.add_heading("7.1 Sensibilidad del liderazgo de la Línea B", level=2)
    add_para(
        doc,
        "Se comparó el reparto por línea de 2025 laborable usando todos los slots observados y aplicando exclusivamente la ventana habitual. El objetivo fue verificar si las extensiones nocturnas o eventos especiales explicaban el liderazgo de la Línea B.",
    )
    add_table(
        doc,
        ["Línea", "Share sin ajuste", "Share habitual", "Volumen excluido", "Ranking habitual"],
        [
            ["B", "23,52%", "23,56%", "0,46%", "1"],
            ["D", "22,11%", "21,98%", "1,22%", "2"],
            ["A", "19,57%", "19,68%", "0,10%", "3"],
            ["C", "13,84%", "13,82%", "0,76%", "4"],
            ["H", "10,97%", "10,92%", "1,16%", "5"],
            ["E", "10,00%", "10,04%", "0,17%", "6"],
        ],
        [1000, 1900, 1900, 1900, 2660],
        font_size=8.8,
    )
    add_para(
        doc,
        "Conclusión: la Línea B permanece primera y la Línea D segunda. La exclusión de las franjas fuera del horario habitual reduce menos de medio punto porcentual del volumen horario de la B; por lo tanto, su liderazgo corresponde al uso habitual y no a extensiones nocturnas. El notebook principal conserva el cálculo, el gráfico comparativo y la exportación CSV para reutilizarlos en la presentación.",
        italic=True,
        color=MUTED,
    )

    prototipos = read_optional_csv("prototipos_maxvol.csv")
    if not prototipos.empty:
        doc.add_heading("7.2 Prototipos MaxVol encontrados", level=2)
        ventanas = {
            "share_06_10": "Mañana",
            "share_10_16": "Mediodía",
            "share_16_21": "Tarde",
            "share_21_24": "Noche",
        }
        proto_rows = []
        for r in prototipos.sort_values(["segmento_dia", "orden_maxvol"]).itertuples(index=False):
            props = {c: float(getattr(r, c)) for c in ventanas}
            dominante = max(props, key=props.get)
            condicion = "Referencia" if bool(r.referencia_modelo) else "Fuera de referencia"
            if bool(r.anio_parcial):
                condicion += "; parcial"
            proto_rows.append([
                r.prototipo_maxvol,
                r.estacion_id,
                int(r.anio),
                r.hora_pico,
                f"{ventanas[dominante]} {props[dominante]:.1%}",
                condicion,
            ])
        add_table(
            doc,
            ["ID", "Estación", "Año", "Pico", "Ventana dominante", "Condición"],
            proto_rows,
            [1000, 3000, 800, 1000, 1760, 1800],
            font_size=8.1,
        )
        fuera = int((~prototipos["referencia_modelo"].astype(bool)).sum())
        add_para(
            doc,
            f"Resultado clave: {fuera} de {len(prototipos)} prototipos pertenecen a perfiles fuera de la referencia estable (pandemia o año parcial). Esto confirma que MaxVol sobre todos los datos captura rupturas temporales además de patrones estructurales. Conviene conservar este resultado global y agregar una segunda selección restringida a la referencia para comparar prototipos estables frente a excepcionales.",
            italic=True,
            color=MUTED,
        )
        add_bullet(doc, "Laborable: Catalinas 2026 representa una concentración vespertina extrema; Plaza de los Virreyes 2015, una concentración matutina extrema; Independencia 2019 aporta una forma de mediodía con mayor extensión nocturna.")
        add_bullet(doc, "Sábado: San Pedrito 2021 concentra el mediodía; Constitución 2020 desplaza peso a la mañana; Facultad de Derecho 2021 concentra la tarde; Medrano 2014 representa una forma más amplia de mediodía.")
        add_bullet(doc, "Domingo/Feriado: Constitución y Juramento 2020, Facultad de Derecho 2018 y Pasco 2021 cubren desde perfiles amplios de tarde hasta un extremo nocturno con pico a las 21:00.")

    doc.add_heading("8. Archivos y resultados principales", level=1)
    add_table(
        doc,
        ["Archivo", "Contenido"],
        [
            ["auditoria_por_anio.csv", "Fechas, intervalos, cobertura y estaciones por año"],
            ["auditoria_estaciones_canonicas.csv", "Identificadores crudos, observados canónicos e inventario esperado"],
            ["alias_estaciones_canonicas.csv", "Textos originales que convergen en una misma línea-estación canónica"],
            ["estaciones_descartadas.csv", "Registros excluidos con motivo, station-slots y pasajeros"],
            ["scores_fpca_clusters_segmentados.csv", "Scores, clusters, distancias y metadatos por curva"],
            ["perfiles_suavizados_segmentados.csv", "Curvas funcionales suavizadas"],
            ["caracterizacion_clusters.csv", "Picos, ventanas horarias, volumen y silhouette"],
            ["prototipos_maxvol.csv", "Prototipos y rasgos interpretables"],
            ["diagnostico_maxvol.csv", "Dimensión, iteraciones, coeficiente final y log-volumen"],
            ["interpretacion_clusters_maxvol.md", "Narrativa automática de grupos y prototipos"],
            ["mapa_estaciones_perfil_tipico.csv", "Perfiles mediana estación-año-segmento-slot, listos para el mapa"],
            ["mapa_estaciones_georreferenciadas.csv", "Capa canónica de coordenadas y procedencia oficial"],
            ["resumen_capa_mapa_dashboard.csv", "Cobertura y tamaño de la capa geográfica por año y segmento"],
            ["analisis_liderazgo_linea_b_horario_habitual.csv", "Sensibilidad del ranking por línea al alcance horario corriente"],
            ["creacion_dashboard.ipynb", "Aplicación Dash con mapa, filtros, animación y linking and brushing"],
            ["resultados_proyecto_subte.zip", "Paquete consolidado de resultados"],
        ],
        [3600, 5760],
        font_size=9,
    )

    doc.add_heading("9. Registro de decisiones", level=1)
    add_table(
        doc,
        ["Fecha", "Decisión", "Justificación"],
        [
            ["2026-08-31", "Usar formato_fechas_por_archivo.csv", "Consistencia histórica y eliminación de ambigüedad día/mes"],
            ["2026-08-31", "Parseo por bloques para 2020", "El mismo archivo cambia de DD/MM a MM/DD y vuelve a DD/MM"],
            ["2026-08-31", "Parseo mixto ISO/DD-MM para 2021", "Los meses contienen más de una representación"],
            ["2026-08-31", "Excluir PM", "El Premetro no pertenece al universo analítico solicitado"],
            ["2026-08-31", "Normalizar estaciones sobre station-slot", "Corregir conteos sin regenerar la caché cruda"],
            ["2026-08-31", "Separar observado de esperado", "Una estación sin registros no equivale a una estación inexistente"],
            ["2026-08-31", "Agregar MaxVol afín por segmento", "Seleccionar observaciones reales que representen la diversidad del espacio FPCA"],
            ["2026-08-31", "Preparar pulso geográfico interactivo", "Unir perfiles típicos a 90 coordenadas oficiales y distinguir volumen absoluto de intensidad relativa"],
            ["2026-08-31", "Mantener bitácora viva", "Trazabilidad reproducible de decisiones y resultados"],
            ["2026-09-17", "Usar operación habitual como alcance del dashboard", "Excluir extensiones por eventos y observaciones aisladas que no describen el día típico"],
            ["2026-09-17", "Conservar todos los datos dentro del horario habitual", "El alcance se define exclusivamente por las ventanas operativas, sin umbral adicional por franja"],
            ["2026-09-17", "Fijar la escala de burbujas entre horas", "Impedir que valores mínimos parezcan grandes por reescalado cuadro a cuadro"],
            ["2026-09-17", "Persistir selecciones por gráfico", "Hacer estable el linking and brushing y permitir reset selectivo desde la barra de Plotly"],
            ["2026-09-17", "Validar liderazgo de la Línea B", "El ranking se mantiene al restringir el análisis a horarios habituales"],
        ],
        [1500, 3300, 4560],
        font_size=8.5,
    )

    doc.add_heading("10. Próximos pasos", level=1)
    add_para(doc, "Completado en esta versión", bold=True, color=DARK_BLUE)
    add_bullet(doc, "Ejecución completa sobre la caché sin releer las 122 fuentes.")
    add_bullet(doc, "Conteos canónicos confirmados y descartes exportados con motivo.")
    add_bullet(doc, "FPCA y clusters recalculados con identidad de estación corregida.")
    add_bullet(doc, "Once prototipos MaxVol globales seleccionados, graficados e interpretados.")
    add_bullet(doc, "Dashboard Dash final con mapa oficial, animación horaria, identidad visual y linking and brushing persistente.")
    add_bullet(doc, "Alcance de operación habitual aplicado y liderazgo de la Línea B validado mediante análisis de sensibilidad.")
    add_para(doc, "Pendiente", bold=True, color=DARK_BLUE)
    next_steps = [
        "Validar los 21 días faltantes informados para 2021 contra los archivos fuente y distinguir ausencia real de error de parseo.",
        "Inspeccionar los intervalos no estándar mediante ejemplos concretos antes de decidir si corregirlos o excluirlos.",
        "Evaluar sensibilidad de MaxVol excluyendo pandemia y años parciales; comparar prototipos globales y de referencia.",
        "Evaluar un modo opcional, separado y apagado por defecto para explorar servicios extraordinarios y eventos especiales.",
        "Incorporar cierres temporales y aperturas de estaciones a un inventario histórico por fecha, no solo por año.",
        "Actualizar este documento ante cada nueva regla, cambio de caché, decisión metodológica o resultado estable.",
    ]
    for step in next_steps:
        add_number(doc, step)

    doc.add_heading("11. Procedimiento de actualización", level=1)
    add_number(doc, "Modificar el notebook o la regla de limpieza y asignar una versión explícita.")
    add_number(doc, "Ejecutar solamente desde la capa mínima necesaria y conservar la caché no afectada.")
    add_number(doc, "Validar conteos, descartes, fechas, intervalos y resultados estadísticos.")
    add_number(doc, "Agregar una fila al Registro de decisiones y actualizar Estado, Resultados y Próximos pasos.")
    add_number(doc, "Regenerar este DOCX ejecutando subte_fpca_work/documentacion/actualizar_bitacora.py y revisar su render.")

    doc.add_heading("12. Fuentes de referencia", level=1)
    sources = [
        ("Diapositivas de la clase: reducción dimensional, FPCA y MaxVol", "https://rballester.github.io/uba_winter_school/#/dimensionality-reduction"),
        ("Buenos Aires Data: estaciones de Subte", "https://data.buenosaires.gob.ar/dataset/subte-estaciones/resource/juqdkmgo-1992-resource"),
        ("Historia oficial de la Línea H", "https://buenosaires.gob.ar/gcaba_historico/jefaturadegabinete/movilidad/subte/nuestra-historia/nuestra-historia/linea-h"),
        ("Historia oficial de la red de Subte", "https://buenosaires.gob.ar/gcaba_historico/jefaturadegabinete/movilidad/subte/nuestra-historia"),
        ("Información oficial sobre reapertura de Plaza de Mayo", "https://buenosaires.gob.ar/gcaba_historico/noticias/vuelven-estar-operativas-15-estaciones-de-subte"),
        ("Horarios oficiales del servicio de Subte", "https://buenosaires.gob.ar/gcaba_historico/infraestructura/subte/horarios-del-servicio"),
        ("Ejemplo oficial de horario extendido de la Línea B", "https://buenosaires.gob.ar/gcaba_historico/noticias/la-linea-b-extiende-su-horario-tambien-los-sabados"),
    ]
    for label, url in sources:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        add_hyperlink(p, label, url)

    add_para(
        doc,
        "Control de mantenimiento: este documento describe el estado validado del proyecto al cierre de la fecha indicada. Debe regenerarse después de cada ejecución que cambie reglas, datos, clusters o prototipos.",
        italic=True,
        color=MUTED,
        after=0,
    )

    DOC_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOC_PATH)
    print(DOC_PATH)


if __name__ == "__main__":
    build()
