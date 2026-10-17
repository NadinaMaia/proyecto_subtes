"""Genera el guion oral y propuesta de dashboard del TP final."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from actualizar_bitacora import (
    BLUE,
    DARK_BLUE,
    GOLD,
    INK,
    LIGHT_BLUE,
    MUTED,
    add_bullet,
    add_hyperlink,
    add_number,
    add_page_number,
    add_para,
    add_table,
    configure_styles,
    OxmlElement,
    qn,
    set_cell_shading,
    set_table_geometry,
    style_run,
)


ROOT = Path(__file__).resolve().parents[2]
DOC_PATH = ROOT / "subte_fpca_work" / "documentacion" / "presentacion_proyecto.docx"


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def build() -> None:
    doc = Document()
    configure_styles(doc)

    # Portada editorial sobria.
    portada = doc.add_paragraph()
    portada.alignment = WD_ALIGN_PARAGRAPH.CENTER
    portada.paragraph_format.space_before = Pt(120)
    portada.paragraph_format.space_after = Pt(16)
    style_run(portada.add_run("PRESENTACIÓN DEL PROYECTO"), size=12, bold=True, color=BLUE)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(12)
    style_run(titulo.add_run("El pulso espacio-temporal del\nSubte de Buenos Aires"), size=26, bold=True, color=DARK_BLUE)

    subtitulo = doc.add_paragraph()
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitulo.paragraph_format.space_after = Pt(28)
    style_run(
        subtitulo.add_run(
            "FPCA, clustering y MaxVol para describir patrones horarios, "
            "prototipos y cambios de la red (2014–2026)"
        ),
        size=12, color=MUTED, italic=True,
    )

    ficha = doc.add_table(rows=2, cols=2)
    ficha.style = "Table Grid"
    set_table_geometry(ficha, [4680, 4680])
    for r in ficha.rows:
        for c in r.cells:
            set_cell_shading(c, LIGHT_BLUE)
    ficha.cell(0, 0).text = "Producto final"
    ficha.cell(0, 1).text = "Dashboard interactivo"
    ficha.cell(1, 0).text = "Fecha de actualización"
    ficha.cell(1, 1).text = date.today().isoformat()
    for row in ficha.rows:
        for col_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    style_run(run, size=10, bold=(col_idx == 0), color=DARK_BLUE if col_idx == 0 else INK)
    header_pr = ficha.rows[0]._tr.get_or_add_trPr()
    header_mark = OxmlElement("w:tblHeader")
    header_mark.set(qn("w:val"), "true")
    header_pr.append(header_mark)

    doc.add_page_break()

    header = doc.sections[0].header.paragraphs[0]
    header.text = "TP final · FPCA del Subte de Buenos Aires"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        style_run(run, size=9, color=MUTED)
    footer = doc.sections[0].footer.paragraphs[0]
    add_page_number(footer)

    heading(doc, "1. Mensaje central para abrir la presentación")
    p = add_para(doc, after=10)
    style_run(p.add_run("Versión oral (30 segundos). "), bold=True, color=DARK_BLUE)
    style_run(
        p.add_run(
            "Este trabajo convierte los registros de molinetes en una lectura de la red: "
            "muestra cómo la intensidad de uso se desplaza entre estaciones durante el día, "
            "cómo esos perfiles horarios se agrupan y cómo cambian en la historia reciente. "
            "El dashboard no busca solo mostrar pasajeros: busca explicar la heterogeneidad "
            "espacial y temporal de la demanda de Subte."
        )
    )

    heading(doc, "2. Problema, objetivo y pregunta guía")
    heading(doc, "Problema", level=2)
    add_para(
        doc,
        "La red no se usa de manera homogénea. Una misma cantidad diaria de pasajeros puede "
        "ocultar patrones muy diferentes: estaciones con pico de ingreso matutino, nodos de "
        "transbordo activos durante todo el día y estaciones de actividad vespertina. Además, "
        "la red se amplió y atravesó una ruptura extraordinaria durante la pandemia."
    )
    heading(doc, "Objetivo general", level=2)
    p = add_para(doc, after=10)
    style_run(
        p.add_run(
            "Construir un dashboard analítico, reproducible y georreferenciado que permita "
            "explorar cómo se distribuye y evoluciona la demanda de las estaciones durante el día "
            "y entre años, distinguiendo tipo de día, cambios de red y discontinuidades temporales."
        ),
        bold=True, color=DARK_BLUE,
    )
    heading(doc, "Pregunta guía", level=2)
    add_para(
        doc,
        "¿Cómo cambió la geografía y la temporalidad de la demanda de la red de Subte de Buenos Aires, "
        "y qué perfiles de uso caracterizan a las estaciones en la operación reciente?",
        bold=True,
    )
    heading(doc, "Objetivos específicos", level=2)
    for item in [
        "Visualizar el pulso horario de las estaciones sobre un mapa, con intervalos de 15 minutos.",
        "Comparar patrones por día laborable, sábado y domingo/feriado sin mezclar comportamientos distintos.",
        "Resumir curvas horarias complejas mediante FPCA y encontrar perfiles similares mediante clustering.",
        "Seleccionar con MaxVol prototipos reales y diversos que ayuden a interpretar los clusters.",
        "Separar cambios estructurales de red, cobertura y pandemia de los patrones operativos actuales.",
    ]:
        add_bullet(doc, item)

    heading(doc, "3. Stakeholders y valor del producto")
    add_table(
        doc,
        ["Stakeholder", "Necesidad", "Cómo aporta el dashboard"],
        [
            ["Docente y jurado", "Rigor metodológico y una historia clara.", "Expone trazabilidad de limpieza, cobertura, FPCA, clusters y prototipos."],
            ["SBASE / operación", "Identificar concentración horaria y espacial de la demanda.", "Permite detectar dónde y cuándo la demanda típica crece, y comparar perfiles de estación."],
            ["Gobierno de la Ciudad / planificación", "Entender cambios territoriales y de red.", "Aporta una lectura histórica contextualizada para intervenciones, accesibilidad e inversión."],
            ["Personas usuarias y público", "Comprender el comportamiento de la red.", "Ofrece una visualización intuitiva del movimiento de la demanda a lo largo del día."],
            ["Equipo analítico", "Reproducibilidad y preguntas futuras.", "Deja datos agregados, indicadores de calidad y supuestos visibles para futuras extensiones."],
        ],
        [1500, 2800, 5060],
        font_size=8.7,
    )
    add_para(
        doc,
        "Alcance responsable: el tablero aporta evidencia descriptiva y exploratoria. No reemplaza "
        "datos de capacidad, tiempos de viaje ni estudios causales necesarios para decidir una intervención operativa.",
        italic=True, color=MUTED,
    )

    heading(doc, "4. Tres narrativas posibles para el dashboard")
    add_table(
        doc,
        ["Enfoque", "Qué cuenta", "Evaluación"],
        [
            ["Historia completa", "Evolución 2014–2026, aperturas, pandemia y recuperación.", "Valioso para contexto, pero puede diluir la lectura operativa del día a día."],
            ["Pulso actual", "Mapa animado del año más reciente y perfiles recientes.", "Muy intuitivo, pero sin contexto puede confundir cambios de red o un año parcial con una tendencia."],
            ["Híbrido: red + pulso", "Primero contexto histórico; luego comportamiento actual y perfiles FPCA.", "Recomendado: integra la historia sin perder una experiencia interactiva focalizada."],
        ],
        [1650, 3870, 3840],
        font_size=9,
    )
    p = add_para(doc, after=8)
    style_run(p.add_run("Decisión recomendada. "), bold=True, color=DARK_BLUE)
    style_run(
        p.add_run(
            "Usar una narrativa híbrida: “de la evolución de la red al pulso actual”. "
            "El mapa animado es la puerta de entrada; la historia 2014–2026 explica por qué "
            "ese patrón no debe leerse fuera de contexto; FPCA, clusters y MaxVol dan profundidad interpretativa."
        )
    )

    heading(doc, "5. Arquitectura propuesta del dashboard")
    add_table(
        doc,
        ["Módulo", "Pregunta que responde", "Elementos principales"],
        [
            ["1. Pulso de la red", "¿Dónde se concentra la demanda a cada hora?", "Mapa animado, selector de año, segmento de día y deslizador de 15 minutos."],
            ["2. Evolución", "¿Qué cambió entre años y qué parte es comparable?", "Serie anual, cobertura, estaciones activas, hitos de red y comparador de años."],
            ["3. Perfiles", "¿Qué tipos de comportamiento horario existen?", "Curvas FPCA, clusters por segmento y mapa/ranking de estaciones por perfil."],
            ["4. Prototipos", "¿Qué casos reales representan mejor la diversidad?", "Tarjetas MaxVol, curva del prototipo, cobertura y condición temporal."],
        ],
        [1700, 3200, 4460],
        font_size=8.9,
    )

    heading(doc, "6. Diseño del mapa interactivo")
    add_para(
        doc,
        "La unidad de visualización es un perfil de “día típico”: la mediana de pasajeros en cada "
        "intervalo de 15 minutos, usando solamente días con cobertura suficiente. La mediana hace "
        "más robusta la animación ante eventos o lecturas anómalas de una fecha puntual."
    )
    for item in [
        "Punto por estación georreferenciada. Tamaño = pasajeros típicos del intervalo: magnitud absoluta.",
        "Color/saturación = porcentaje del total diario típico de esa estación: intensidad relativa y momento de pico.",
        "Reproducción automática y deslizador horario; hover con línea, nombre, volumen típico, porcentaje y cantidad de días incluidos.",
        "Filtros mínimos: año, laborable/sábado/domingo-feriado, línea y opción de métrica absoluta o relativa.",
        "Advertencia visible cuando el año sea parcial o tenga cobertura insuficiente.",
    ]:
        add_bullet(doc, item)

    p = add_para(doc, after=8)
    style_run(p.add_run("Lectura sugerida para explicar el mapa. "), bold=True, color=DARK_BLUE)
    style_run(
        p.add_run(
            "“Los círculos grandes muestran dónde hay más validaciones en ese intervalo. "
            "El color muestra si ese momento representa una parte alta del día propio de la estación. "
            "Así distinguimos una estación grande de una estación que está justo en su momento de mayor actividad.”"
        )
    )

    heading(doc, "¿Usar solo 2026?", level=2)
    add_para(
        doc,
        "No como única vista. 2026 es útil para responder “cómo se comporta hoy la red”, pero es parcial y no representa "
        "un año cerrado. Conviene ofrecerlo como modo actual, acompañado por una etiqueta de cobertura. Para una referencia "
        "estable, el tablero puede sugerir 2025 —último año completo— con el mismo tipo de día y horario."
    )
    heading(doc, "¿Para qué sirve 2014–2026?", level=2)
    add_para(
        doc,
        "La serie histórica no debe reproducirse como una única película continua: cambió la red, se abrieron estaciones y "
        "la pandemia alteró radicalmente el uso. Su valor es explicar la transformación estructural, comparar años completos "
        "bajo filtros homogéneos y mostrar si un perfil actual es nuevo o persistente. Las comparaciones históricas deben usar "
        "valores relativos o estaciones comunes y exhibir cobertura e hitos de red."
    )

    heading(doc, "7. Cómo entra FPCA, clustering y MaxVol en la historia")
    add_para(
        doc,
        "El mapa responde “dónde y cuándo”. FPCA y clustering responden “qué formas de uso se repiten”. "
        "Las curvas de estación–año–segmento se normalizan en forma para comparar la distribución horaria, no solo el volumen. "
        "Luego, MaxVol elige perfiles reales, diversos y extremos dentro del espacio de scores: son ejemplos interpretables, "
        "no promedios artificiales ni clusters adicionales."
    )
    for item in [
        "Comparar clusters únicamente dentro del mismo segmento de día.",
        "Usar MaxVol para contar con casos concretos que ilustren la diversidad de perfiles observada.",
        "Señalar los prototipos de 2020–2021 y de 2026 como casos que pueden reflejar una ruptura temporal o cobertura parcial.",
    ]:
        add_bullet(doc, item)

    heading(doc, "8. Guion oral sugerido (8–10 minutos)")
    add_table(
        doc,
        ["Tiempo", "Momento", "Texto/acción sugerida"],
        [
            ["0:00–0:45", "Apertura", "Usar el mensaje central. Mostrar una captura del mapa y plantear que el desafío es ver la red como un sistema que cambia durante el día."],
            ["0:45–1:30", "Pregunta y usuarios", "Presentar la pregunta guía y aclarar que el tablero sirve para exploración, planificación y evaluación académica."],
            ["1:30–2:45", "Datos y calidad", "Explicar brevemente fechas mixtas, intervalos de 15 minutos, exclusión del Premetro, duplicados y normalización de estaciones. Mostrar cobertura como condición de confianza."],
            ["2:45–4:30", "Pulso horario", "Animar el mapa para un laborable. Explicar tamaño absoluto, color relativo y por qué se usa un día típico mediano."],
            ["4:30–5:45", "Historia", "Pasar a evolución anual. Mostrar que 2014–2026 contextualiza aperturas, pandemia y recuperación; no es una comparación ingenua de todos los años."],
            ["5:45–7:30", "FPCA y clusters", "Mostrar las curvas/profiles por segmento de día. Traducir los grupos a horas de pico y concentración, sin sobreinterpretar causalidad."],
            ["7:30–8:45", "MaxVol", "Presentar prototipos reales: son casos frontera que representan diversidad y ayudan a leer los clusters."],
            ["8:45–9:30", "Cierre", "Concluir que el dashboard vincula geografía, tiempo e historia. Mencionar límites y próximos pasos."],
        ],
        [900, 1850, 6610],
        font_size=8.55,
    )

    heading(doc, "9. Frases de transición y cierre")
    for item in [
        "Del mapa a la historia: “Antes de comparar años, necesitamos separar cambios reales de la demanda de cambios en la red y en la cobertura.”",
        "De la historia a FPCA: “La evolución nos muestra cuánto y dónde cambió; FPCA nos permite resumir cómo cambió la forma del día.”",
        "De clusters a MaxVol: “El cluster es un patrón promedio; MaxVol nos devuelve ejemplos reales y diversos para ponerle cara a esa variación.”",
        "Cierre: “La contribución del proyecto es hacer visible el pulso de la red y, al mismo tiempo, dar herramientas para interpretarlo con cuidado histórico y estadístico.”",
    ]:
        add_bullet(doc, item)

    heading(doc, "10. Próximos pasos antes de la entrega")
    for item in [
        "Construir la interfaz final con los cuatro módulos priorizados; si el tiempo es limitado, empezar por Pulso, Evolución y Perfiles.",
        "Incorporar controles de cobertura y leyendas metodológicas visibles en cada vista.",
        "Agregar un comparador de dos años completos con métricas relativas y universo de estaciones comunes.",
        "Probar el dashboard con una historia guiada de tres hallazgos y documentar cualquier nueva limpieza en la bitácora técnica.",
    ]:
        add_number(doc, item)

    heading(doc, "Fuentes de referencia")
    p = add_para(doc, "Coordenadas de estaciones: ")
    add_hyperlink(p, "Buenos Aires Data — Subte: Estaciones", "https://data.buenosaires.gob.ar/dataset/subte-estaciones")
    add_para(doc, "Metodología de la cursada: ")
    p = doc.paragraphs[-1]
    add_hyperlink(p, "Diapositivas de FPCA y MaxVol", "https://rballester.github.io/uba_winter_school/#/dimensionality-reduction")

    DOC_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOC_PATH)
    print(f"Documento creado: {DOC_PATH}")


if __name__ == "__main__":
    build()
